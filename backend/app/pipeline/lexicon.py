"""Excel word lists ang gamit sa health, question at opinion checks."""
import re
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook

from .. import config

SHEET_HEALTH, SHEET_QUESTIONS, SHEET_OPINIONS = "health_terms", "questions", "opinions"
_SENTENCE_SPLIT = re.compile(r"(?<=[.!?])\s+")
MIN_CLAIM_WORDS = 3


class LexiconError(RuntimeError):
    pass


def _rows(path: Path, sheet: str, columns: list[str]) -> list[dict]:
    """Basahin ang sheet; laktawan ang blank at comment rows."""
    if not path.is_file():
        raise LexiconError(f"Lexicon file not found: {path}")
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise LexiconError(f"{path}: sheet '{sheet}' not found (sheets: {', '.join(wb.sheetnames)})")
        it = wb[sheet].iter_rows(values_only=True)
        header = [str(c).strip().lower() if c is not None else "" for c in next(it, ())]
        missing = [c for c in columns if c not in header]
        if missing:
            raise LexiconError(f"{path}, sheet '{sheet}': missing column(s) {', '.join(missing)}")
        index = {c: header.index(c) for c in header if c}
        out = []
        for n, row in enumerate(it, start=2):
            cells = {c: ("" if i >= len(row) or row[i] is None else str(row[i]).strip()) for c, i in index.items()}
            first = cells[columns[0]]
            if first and not first.startswith("#"):
                cells["_line"] = n
                out.append(cells)
        return out
    finally:
        wb.close()


def _phrase(text: str) -> str:
    """Gawing regex ang phrase na flexible sa spaces."""
    return r"\s+".join(re.escape(w) for w in text.lower().split())


# I-check kung allowed ang match type ng row.
def _check_match(path: Path, sheet: str, row: dict, allowed: tuple[str, ...]) -> str:
    match = (row.get("match") or allowed[0]).lower()
    if match not in allowed:
        raise LexiconError(f"{path}, sheet '{sheet}', row {row['_line']}: match must be {' or '.join(allowed)}, got '{match}'")
    return match


def load_health_patterns(path: Path) -> list[str]:
    """Buuin ang exact/prefix patterns mula sa health terms."""
    patterns = []
    for row in _rows(path, SHEET_HEALTH, ["term", "match"]):
        match = _check_match(path, SHEET_HEALTH, row, ("exact", "prefix"))
        patterns.append(_phrase(row["term"]) + (r"\w*" if match == "prefix" else ""))
    if not patterns:
        raise LexiconError(f"No health terms in {path}")
    return patterns


def load_phrases(path: Path, sheet: str, allowed: tuple[str, ...]) -> dict[str, list[str]]:
    """Igrupo ang phrase patterns ayon sa match type."""
    grouped: dict[str, list[str]] = {m: [] for m in allowed}
    for row in _rows(path, sheet, ["phrase", "match"]):
        grouped[_check_match(path, sheet, row, allowed)].append(_phrase(row["phrase"]))
    return grouped


# Buuin at i-cache ang health, question at opinion rules.
@lru_cache(maxsize=1)
def _compiled() -> dict[str, re.Pattern | None]:
    path = config.LEXICON_PATH
    health = sorted(load_health_patterns(path), key=len, reverse=True)
    questions = load_phrases(path, SHEET_QUESTIONS, ("start", "end", "contains"))
    opinions = load_phrases(path, SHEET_OPINIONS, ("start", "contains"))

    # Pagsamahin ang phrases sa regex; None kung walang entries.
    def pick(parts: list[str], template: str) -> re.Pattern | None:
        return re.compile(template.format("|".join(sorted(parts, key=len, reverse=True))), re.IGNORECASE) if parts else None

    return {
        "health": re.compile(r"\b(?:" + "|".join(health) + r")\b", re.IGNORECASE),
        "question_start": pick(questions["start"], r"^\s*(?:{})\b"),
        "question_end": pick(questions["end"], r"\b(?:{})\s*[?.!]*$"),
        "question_contains": pick(questions["contains"], r"\b(?:{})\b"),
        "opinion_start": pick(opinions["start"], r"^\s*(?:{})\b"),
        "opinion_contains": pick(opinions["contains"], r"\b(?:{})\b"),
    }


def reload_lexicon() -> None:
    """I-clear ang cache para mabasa ulit ang edited workbook."""
    _compiled.cache_clear()


# I-check kung may health term ang text.
def is_health_text(text: str) -> bool:
    return bool(_compiled()["health"].search(text))


_WHO_ORG = re.compile(r"^\s*WHO\b")  # WHO organization ito, hindi ang question word na who.


def is_question(sentence: str) -> bool:
    """I-check kung tanong; huwag ituring na question word ang WHO."""
    s = sentence.strip()
    if s.endswith("?"):
        return True
    c = _compiled()
    if c["question_start"] and c["question_start"].search(s) and not _WHO_ORG.match(s):
        return True
    return any(c[k] and c[k].search(s) for k in ("question_end", "question_contains"))


# I-check kung may personal opinion o feeling.
def is_subjective(sentence: str) -> bool:
    c = _compiled()
    return bool((c["opinion_start"] and c["opinion_start"].match(sentence))
                or (c["opinion_contains"] and c["opinion_contains"].search(sentence)))


# Ayusin ang spaces at hatiin ang text sa sentences.
def sentences(text: str) -> list[str]:
    return [s.strip() for s in _SENTENCE_SPLIT.split(" ".join(text.split())) if s.strip()]


def is_claim_like(sentence: str) -> bool:
    """Hanapin ang health assertion na hindi tanong o opinion."""
    return (
        len(sentence.split()) >= MIN_CLAIM_WORDS
        and is_health_text(sentence)
        and not is_question(sentence)
        and not is_subjective(sentence)
    )
